"""One-time migration: local SQLite + the legacy Excel recipe/ECR workbooks
-> a Postgres database (e.g. Supabase).

Run it once, from the project root, with the TARGET database URL set:

    # PowerShell
    $env:DATABASE_URL = "postgresql://USER:PASS@HOST:5432/postgres?sslmode=require"
    python "Main Codes/migrate_to_supabase.py"

It reads:
  - Database/plant.db                 (equipment, products, users, orders,
                                        batches, allocations, cleaning steps,
                                        audit log)
  - Database/recipes.xlsx             (product recipes -> recipe_* tables)
  - Database/ecr_templates.xlsx       (cleaning templates -> ecr_template_* tables)
and writes all of it into DATABASE_URL. Source files are never modified.

The target must be empty (no users). Re-run with --force to wipe the managed
tables on the target first.
"""
import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from sqlalchemy import create_engine, func, insert, select, text  # noqa: E402

from batch_planner import config  # noqa: E402
from batch_planner.models import (  # noqa: E402
    Allocation, AuditLog, Base, Batch, CleaningStep, EcrTemplate, EcrTemplateStep,
    Equipment, Order, Product, RecipeSheet, RecipeStage, User,
)

# SQLite source tables, in an order that satisfies every foreign key.
CORE_MODELS = [Equipment, Product, User, Order, Batch, Allocation, CleaningStep, AuditLog]
# Target tables cleared by --force (reverse dependency order) and whose
# Postgres id sequences need bumping after an id-preserving copy.
ALL_MODELS = CORE_MODELS + [RecipeSheet, RecipeStage, EcrTemplate, EcrTemplateStep]
SEQ_MODELS = [Order, Batch, Allocation, CleaningStep, AuditLog, RecipeStage, EcrTemplateStep]


def _copy_table(model, src_conn, tgt_conn) -> int:
    table = model.__table__
    rows = [dict(r._mapping) for r in src_conn.execute(select(table))]
    if rows:
        tgt_conn.execute(insert(table), rows)
    return len(rows)


def _read_recipes_xlsx(path: Path) -> dict[str, list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    out: dict[str, list[dict]] = {}
    for name in wb.sheetnames:
        if name.lower() == "readme":
            continue
        ws = wb[name]
        stages: list[dict] = []
        col = 2
        while True:
            stage_name = ws.cell(row=1, column=col).value
            if stage_name in (None, ""):
                break
            stages.append({
                "sheet": name, "seq": col - 1, "name": str(stage_name),
                "op_minutes": float(ws.cell(row=2, column=col).value or 0),
                "clean_minutes": float(ws.cell(row=3, column=col).value or 0),
                "equipment_ids": str(ws.cell(row=4, column=col).value or ""),
                "temperature": str(ws.cell(row=5, column=col).value or ""),
                "actual_temperature": str(ws.cell(row=6, column=col).value or ""),
            })
            col += 1
        out[name] = stages
    wb.close()
    return out


def _read_ecr_templates_xlsx(path: Path) -> dict[str, list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    out: dict[str, list[dict]] = {}
    for name in wb.sheetnames:
        if name.lower() == "readme":
            continue
        ws = wb[name]
        steps: list[dict] = []
        col = 2
        while True:
            step_name = ws.cell(row=1, column=col).value
            if step_name in (None, ""):
                break
            steps.append({
                "template_key": name, "seq": col - 1, "name": str(step_name),
                "op_minutes": float(ws.cell(row=2, column=col).value or 0),
            })
            col += 1
        out[name] = steps
    wb.close()
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", default=os.environ.get("DATABASE_URL"),
                        help="Target DB URL (default: $DATABASE_URL)")
    parser.add_argument("--force", action="store_true",
                        help="Wipe the managed tables on the target first")
    args = parser.parse_args()

    if not args.target:
        print("ERROR: set DATABASE_URL (or pass --target) to the Postgres/Supabase URL.")
        return 2
    target_url = config.normalise_db_url(args.target)
    if target_url.startswith("sqlite"):
        print("ERROR: target looks like SQLite. This migrates INTO Postgres.")
        return 2
    if not config.DB_PATH.exists():
        print(f"ERROR: source database not found at {config.DB_PATH}")
        return 2

    src_engine = create_engine(f"sqlite:///{config.DB_PATH}")
    tgt_engine = create_engine(target_url, pool_pre_ping=True)

    print(f"Source : {config.DB_PATH}")
    print(f"Target : {target_url.split('@')[-1]}")
    print("Creating tables on target ...")
    Base.metadata.create_all(tgt_engine)

    with tgt_engine.begin() as tgt, src_engine.connect() as src:
        existing_users = tgt.execute(select(func.count()).select_from(User.__table__)).scalar()
        if existing_users and not args.force:
            print(f"ERROR: target already has {existing_users} user(s). Re-run with --force to overwrite.")
            return 1
        if args.force:
            print("--force: clearing managed tables on target ...")
            for model in reversed(ALL_MODELS):
                tgt.execute(model.__table__.delete())

        print("\nCopying database tables:")
        for model in CORE_MODELS:
            n = _copy_table(model, src, tgt)
            print(f"  {model.__tablename__:<20} {n}")

        recipes_path = config.RECIPES_PATH
        if recipes_path.exists():
            print(f"\nImporting {recipes_path.name}:")
            sheets = _read_recipes_xlsx(recipes_path)
            if sheets:
                tgt.execute(insert(RecipeSheet.__table__),
                            [{"name": s} for s in sheets])
            all_stages = [row for stage_list in sheets.values() for row in stage_list]
            if all_stages:
                tgt.execute(insert(RecipeStage.__table__), all_stages)
            print(f"  recipe_sheets        {len(sheets)}")
            print(f"  recipe_stages        {len(all_stages)}")
        else:
            print(f"\n(skipping {recipes_path.name} — not found)")

        ecr_path = config.ECR_TEMPLATES_PATH
        if ecr_path.exists():
            print(f"\nImporting {ecr_path.name}:")
            templates = _read_ecr_templates_xlsx(ecr_path)
            if templates:
                tgt.execute(insert(EcrTemplate.__table__),
                            [{"key": k} for k in templates])
            all_steps = [row for step_list in templates.values() for row in step_list]
            if all_steps:
                tgt.execute(insert(EcrTemplateStep.__table__), all_steps)
            print(f"  ecr_template_sheets  {len(templates)}")
            print(f"  ecr_template_steps   {len(all_steps)}")
        else:
            print(f"\n(skipping {ecr_path.name} — not found)")

        if tgt_engine.dialect.name == "postgresql":
            print("\nResetting Postgres id sequences ...")
            for model in SEQ_MODELS:
                tn = model.__tablename__
                tgt.execute(text(
                    f"SELECT setval(pg_get_serial_sequence('{tn}', 'id'), "
                    f"GREATEST((SELECT COALESCE(MAX(id), 0) FROM {tn}), 1), true)"
                ))

    print("\nDone. Point the app at DATABASE_URL and sign in with your existing accounts.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
